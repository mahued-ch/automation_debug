import pandas as pd
import sys
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import json

def convert_to_excel(input_file):
    output_file = input_file.rsplit('.', 1)[0] + '.xlsx'
    try:
        df = pd.read_csv(input_file, sep='\t', encoding='utf-16', engine='python', on_bad_lines='skip', skiprows=1)
        
        # Eliminar columnas sin nombre
        df = df.loc[:, ~df.columns.str.contains('^Unnamed')] 
        
        # Guardar en Excel
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Sheet1', index=False)
        
        # Convertir a tabla en Excel
        wb = load_workbook(output_file)
        ws = wb.active
        tabla = Table(displayName="TablaDatos", ref=ws.dimensions)
        estilo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        tabla.tableStyleInfo = estilo
        ws.add_table(tabla)
       
# Ajustar el ancho de las columnas automáticamente
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter  # Obtener la letra de la columna
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 5)  # Añadir un margen adicional
            ws.column_dimensions[column].width = adjusted_width
        
        wb.save(output_file)
                
        # Análisis de columnas específicas
        if 'ID mensaje' in df.columns: # and 'ColumnaAnalizar2' in df.columns:
            print("ID mensaje:")
            print(df['ID mensaje'].describe())
            id_mensaje_counts = df['ID mensaje'].value_counts().to_dict()
            json_output = json.dumps(id_mensaje_counts, ensure_ascii=False, indent=4)
#            print("Resumen de ColumnaAnalizar2:")
#            print(df['ColumnaAnalizar2'].describe())
            print(json_output)

        print(f"Archivo convertido y formateado con éxito: {output_file}")
    except Exception as e:
        print(f"Error al convertir el archivo: {e}")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python script.py <archivo>")
    else:
        convert_to_excel(sys.argv[1])
